from __future__ import annotations

from pathlib import Path

import pytest
from nicegui import ui
from ui_test_support import run_ui, wait_until

from openpyxl import load_workbook

from garmin_data_hub.db.migrate import apply_schema
from garmin_data_hub.db.sqlite import connect_sqlite
from garmin_data_hub.exports.master_export import (
    generate_master_workbook,
    generate_plan_data,
)
from garmin_data_hub.paths import schema_sql_path
from garmin_data_hub.db import queries
from garmin_data_hub.services.plan_persistence import (
    save_generated_plan, get_active_plan_sha256, get_active_plan_snapshot,
)
from garmin_data_hub.ui_nicegui.data import save_planning_settings
from garmin_data_hub.ui_nicegui.data import plan_rows


def _database(tmp_path: Path) -> Path:
    db_path = tmp_path / "garmin.db"
    conn = connect_sqlite(db_path)
    try:
        apply_schema(conn, schema_sql_path())
    finally:
        conn.close()
    return db_path


def _baseline_data():
    return generate_plan_data(
        athlete_name="Offline Athlete",
        age=42,
        lthr=165,
        hrmax=190,
        sodium_mg_per_hr_hot=700,
        event_name="Autumn 10K",
        distance="10K",
        start_date_iso="2026-08-24",
        event_date_iso="2026-09-06",
        run_days_per_week=4,
        long_run_day="Saturday",
    )


@pytest.mark.parametrize("include_workbook", [False, True], ids=["offline", "workbook"])
def test_baseline_page_requires_confirmation_and_refreshes_saved_schedule(
    monkeypatch, tmp_path, include_workbook
):
    db = _database(tmp_path)
    _configure_baseline(db, tmp_path)
    before = get_active_plan_sha256(db)

    async def scenario(user):
        await user.open("/plan")
        label = "Generate baseline + workbook" if include_workbook else "Generate offline baseline"
        user.find(label).click()
        dialog = next(iter(user.find(ui.dialog).elements))
        await wait_until(lambda: dialog.value)
        assert not next(iter(user.find(label).elements)).enabled
        assert get_active_plan_sha256(db) == before
        user.find("Replace range and generate").click()
        await wait_until(lambda: user.notify.contains("Check the replacement acknowledgement first."))
        assert get_active_plan_sha256(db) == before
        user.find("I understand that workouts in this date range will be replaced.").click()
        user.find("Replace range and generate").click()
        await user.should_see("Saved", retries=100)
        rows = plan_rows(db)
        assert rows and rows[-1]["intensity"] == "race"
        assert any(len(grid.options.get("rowData", [])) == len(rows)
                   for grid in user.find(ui.aggrid).elements)
        download = next(element for element in user.client.elements.values()
                        if isinstance(element, ui.button)
                        and element._props.get("label") == "Download last workbook")
        assert download.visible is include_workbook
        assert next(iter(user.find(label).elements)).enabled
        assert (tmp_path / "baseline.xlsx").is_file() is include_workbook
        if include_workbook:
            user.find("Generate offline baseline").click()
            await wait_until(lambda: dialog.value)
            user.find("I understand that workouts in this date range will be replaced.").click()
            user.find("Replace range and generate").click()
            await wait_until(lambda: not dialog.value and
                             next(iter(user.find("Generate offline baseline").elements)).enabled)
            assert not download.visible
        destinations = []
        monkeypatch.setattr(user.navigate, "to", destinations.append)
        user.find("Open Codex Coach").click()
        assert destinations == ["/coach"]

    run_ui(db, scenario)


def test_offline_baseline_persists_rows_consumed_by_plan_schedule(tmp_path):
    db_path = _database(tmp_path)
    inputs, analysis, day_plans, weekly_rows = _baseline_data()

    save_generated_plan(db_path, inputs, analysis, day_plans, weekly_rows)

    schedule = plan_rows(db_path)
    expected_workouts = [day for day in day_plans if day.workout]
    assert len(schedule) == len(expected_workouts)
    assert schedule[0]["date"] == expected_workouts[0].iso_date
    assert schedule[-1]["date"] == expected_workouts[-1].iso_date
    assert schedule[-1]["intensity"] == "race"
    assert "run" in {row["sport"] for row in schedule}
    assert all(row["workout"] for row in schedule)


def test_offline_workbook_contains_generated_calendar(tmp_path):
    output_path = tmp_path / "Offline_Athlete_master_workbook.xlsx"

    result = generate_master_workbook(
        athlete_name="Offline Athlete",
        age=42,
        lthr=165,
        hrmax=190,
        sodium_mg_per_hr_hot=700,
        event_name="Autumn 10K",
        distance="10K",
        start_date_iso="2026-08-24",
        event_date_iso="2026-09-06",
        run_days_per_week=4,
        long_run_day="Saturday",
        out_path=output_path,
    )

    assert result == output_path
    assert output_path.is_file()
    workbook = load_workbook(output_path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == [
            "Calendar",
            "Metrics",
            "Garmin Analysis",
            "Forever Plan",
            "Workout Library",
            "Nutrition",
        ]
        calendar = workbook["Calendar"]
        assert calendar.max_row == 15
        assert calendar["B2"].value == "2026-08-24"
        assert calendar[f"B{calendar.max_row}"].value == "2026-09-06"
    finally:
        workbook.close()


def _configure_baseline(db, output):
    with connect_sqlite(db) as conn:
        queries.set_override_metrics(conn, hrmax=190, lthr=165)
    save_planning_settings(db, {
        "athlete_name": "Offline Athlete", "age": 42, "distance": "10K",
        "event_name": "Autumn 10K", "plan_start": "2026-08-24", "event_date": "2026-09-06",
        "run_days_per_week": 4, "long_run_day": "Saturday", "training_method": "eighty_twenty",
        "sodium_mg_per_hour": 700, "output_directory": str(output),
        "output_filename": "baseline.xlsx",
    })


def test_baseline_page_rejects_a_plan_changed_during_confirmation(tmp_path):
    db = _database(tmp_path)
    _configure_baseline(db, tmp_path)

    async def scenario(user):
        await user.open("/plan")
        user.find("Generate baseline + workbook").click()
        dialog = next(iter(user.find(ui.dialog).elements))
        await wait_until(lambda: dialog.value)
        with connect_sqlite(db) as conn:
            conn.execute("INSERT INTO planned_workout(scheduled_date,workout_name) VALUES (?,?)",
                         ("2026-09-01", "Added from another page"))
            conn.commit()
        before = get_active_plan_snapshot(db)
        user.find("I understand that workouts in this date range will be replaced.").click()
        user.find("Replace range and generate").click()
        await user.should_see("Baseline generation failed", retries=100)
        assert get_active_plan_snapshot(db) == before
        assert not (tmp_path / "baseline.xlsx").exists()
        assert not list(tmp_path.glob(".garmin-data-hub-plan-*.xlsx"))
        assert not next(element for element in user.client.elements.values()
                        if isinstance(element, ui.button)
                        and element._props.get("label") == "Download last workbook").visible
        assert next(iter(user.find("Generate baseline + workbook").elements)).enabled

    run_ui(db, scenario)
